#!/usr/bin/env python3
"""
Genera los archivos de Alborada en los formatos que usa una academia de verdad,
para poder cargarlos en un Proyecto de ChatGPT y sacarle la captura a la placa
"Dónde lo ponés".

  planillas/*.csv  ->  para-cargar/*.xlsx
  documentos/*.md  ->  para-cargar/*.pdf   (vía HTML + Chrome headless)

Uso:  python3 academia-ficticia/_convertir.py
"""

import csv
import html
import re
import subprocess
import sys
from pathlib import Path

RAIZ = Path(__file__).parent
SALIDA = RAIZ / "para-cargar"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

NOMBRES = {
    "alumnos": "Alumnos",
    "pagos": "Pagos",
    "notas-y-asistencia": "Notas y asistencia",
    "reglamento-interno": "Reglamento interno",
    "acuerdo-de-inscripcion": "Acuerdo de inscripción",
    "aranceles-2026": "Aranceles 2026",
    "grilla-horaria-2026": "Grilla horaria 2026",
    "instructivo-secretaria": "Instructivo de secretaría",
}

# ---------------------------------------------------------------- planillas

def csv_a_xlsx(origen: Path, destino: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = destino.stem[:31]

    with origen.open(newline="", encoding="utf-8") as fh:
        filas = list(csv.reader(fh))

    for fila in filas:
        ws.append([_num(v) for v in fila])

    cabecera = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="44546A")
    for celda in ws[1]:
        celda.font = cabecera
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for i, col in enumerate(zip(*filas), start=1):
        ancho = max(len(str(v)) for v in col) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(ancho, 42)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(destino)


def _num(v: str):
    """Los importes y legajos entran como número cuando corresponde."""
    if re.fullmatch(r"-?\d+", v or ""):
        return v if (v.startswith("0") and len(v) > 1) else int(v)
    if re.fullmatch(r"-?\d+\.\d+", v or ""):
        return float(v)
    return v

# ---------------------------------------------------------------- documentos

CSS = """
@page { size: A4; margin: 20mm 18mm; }
body { font: 11pt/1.55 -apple-system, 'Helvetica Neue', Arial, sans-serif; color: #1a1a1a; }
h1 { font-size: 19pt; margin: 0 0 4px; }
h2 { font-size: 14pt; margin: 22px 0 8px; border-bottom: 1px solid #ddd; padding-bottom: 4px; }
h3 { font-size: 12pt; margin: 16px 0 6px; }
p  { margin: 0 0 9px; }
ul { margin: 0 0 10px 18px; padding: 0; }
li { margin-bottom: 4px; }
hr { border: 0; border-top: 1px solid #ddd; margin: 18px 0; }
table { border-collapse: collapse; width: 100%; margin: 10px 0 16px; font-size: 10pt; }
th, td { border: 1px solid #ccc; padding: 5px 8px; text-align: left; }
th { background: #f0f2f5; font-weight: 600; }
strong { font-weight: 600; }
"""


def md_a_html(md: str, titulo: str) -> str:
    out, i = [], 0
    lineas = md.split("\n")

    def inline(t: str) -> str:
        t = html.escape(t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"(?<!\*)\*([^*]+?)\*(?!\*)", r"<em>\1</em>", t)
        t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
        return t

    while i < len(lineas):
        ln = lineas[i]
        if ln.strip().startswith("```"):  # BLOQUE DE CODIGO
            i += 1
            cod = []
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                cod.append(html.escape(lineas[i]))
                i += 1
            i += 1
            out.append("<pre>" + "\n".join(cod) + "</pre>")
            continue
        if ln.strip().startswith(">"):
            cita = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                cita.append(inline(re.sub(r"^\s*>\s?", "", lineas[i])))
                i += 1
            out.append("<blockquote><p>" + " ".join(cita) + "</p></blockquote>")
            continue
        if re.fullmatch(r"\s*", ln):
            i += 1
            continue
        if re.fullmatch(r"-{3,}", ln.strip()):
            out.append("<hr>")
            i += 1
            continue
        m = re.match(r"(#{1,4})\s+(.*)", ln)
        if m:
            n = len(m.group(1))
            out.append(f"<h{n}>{inline(m.group(2))}</h{n}>")
            i += 1
            continue
        if ln.lstrip().startswith("|") and i + 1 < len(lineas) and re.search(r"\|\s*-{2,}", lineas[i + 1]):
            filas = []
            while i < len(lineas) and lineas[i].lstrip().startswith("|"):
                filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            cab, cuerpo = filas[0], filas[2:]
            t = ["<table><thead><tr>"] + [f"<th>{inline(c)}</th>" for c in cab] + ["</tr></thead><tbody>"]
            for f in cuerpo:
                t += ["<tr>"] + [f"<td>{inline(c)}</td>" for c in f] + ["</tr>"]
            out.append("".join(t + ["</tbody></table>"]))
            continue
        if re.match(r"\s*[-*]\s+", ln):
            items = []
            while i < len(lineas) and re.match(r"\s*[-*]\s+", lineas[i]):
                items.append("<li>" + inline(re.sub(r"^\s*[-*]\s+", "", lineas[i])) + "</li>")
                i += 1
            out.append("<ul>" + "".join(items) + "</ul>")
            continue
        parr = []
        while i < len(lineas) and lineas[i].strip() and not re.match(r"\s*([-*]\s|#|\||-{3,})", lineas[i]):
            parr.append(lineas[i].strip())
            i += 1
        if parr:
            out.append(f"<p>{inline(' '.join(parr))}</p>")

    return (
        f"<!doctype html><html lang='es'><head><meta charset='utf-8'>"
        f"<title>{html.escape(titulo)}</title><style>{CSS}</style></head>"
        f"<body>{''.join(out)}</body></html>"
    )


def md_a_pdf(origen: Path, destino: Path) -> None:
    tmp = SALIDA / f".{origen.stem}.html"
    tmp.write_text(md_a_html(origen.read_text(encoding="utf-8"), destino.stem), encoding="utf-8")
    subprocess.run(
        [CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
         f"--print-to-pdf={destino}", f"file://{tmp}"],
        check=True, capture_output=True,
    )
    tmp.unlink()

# ---------------------------------------------------------------- main

def main() -> None:
    SALIDA.mkdir(exist_ok=True)
    hechos = []

    for csvf in sorted((RAIZ / "planillas").glob("*.csv")):
        nombre = NOMBRES.get(csvf.stem, csvf.stem)
        destino = SALIDA / f"{nombre}.xlsx"
        csv_a_xlsx(csvf, destino)
        hechos.append(destino)

    for mdf in sorted((RAIZ / "documentos").glob("*.md")):
        nombre = NOMBRES.get(mdf.stem, mdf.stem)
        destino = SALIDA / f"{nombre}.pdf"
        md_a_pdf(mdf, destino)
        hechos.append(destino)

    print(f"\n{len(hechos)} archivos en {SALIDA}/\n")
    for f in hechos:
        print(f"  {f.name:<32} {f.stat().st_size // 1024} KB")


if __name__ == "__main__":
    if not Path(CHROME).exists():
        sys.exit("No encontré Chrome; hace falta para generar los PDF.")
    main()
